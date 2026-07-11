import json
import os
from datetime import datetime
from pathlib import Path
from typing import Any

from app.core.config import settings


def _to_display_text(text: str) -> str:
    """Reshape Arabic text for PDF rendering."""
    if not text:
        return ""
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display

        reshaped = arabic_reshaper.reshape(str(text))
        return get_display(reshaped)
    except Exception:
        return str(text)


def _find_font_path(font_path: str | None) -> str | None:
    if font_path and Path(font_path).exists():
        return font_path
    # Fallback candidates
    candidates = [
        "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansArabic.ttf",
        "/usr/share/fonts/truetype/freefont/FreeSerif.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for c in candidates:
        if Path(c).exists():
            return c
    return None


def export_pdf(data: dict[str, Any], base_name: str, font_path: str | None = None) -> str:
    from fpdf import FPDF

    report_dir = settings.reports_dir
    report_dir.mkdir(parents=True, exist_ok=True)
    output_path = report_dir / f"{base_name}.pdf"

    class RTLPDF(FPDF):
        def header(self):
            self.set_font("Noto", "", 10)
            self.set_text_color(80, 80, 80)
            self.cell(0, 10, _to_display_text("منصة وكلاء الذكاء الاصطناعي للقطاع غير الربحي"), align="R", new_x="LMARGIN", new_y="NEXT")
            self.ln(2)

        def footer(self):
            self.set_y(-15)
            self.set_font("Noto", "", 8)
            self.set_text_color(128, 128, 128)
            self.cell(0, 10, _to_display_text(f"صفحة {self.page_no()}"), align="C")

    pdf = RTLPDF(orientation="P", unit="mm", format="A4")
    font_file = _find_font_path(font_path)
    if not font_file:
        raise RuntimeError("No suitable font found for Arabic PDF generation")
    pdf.add_font("Noto", "", font_file)
    bold_font = font_file.replace("-Regular", "-Bold")
    if not Path(bold_font).exists():
        bold_font = font_file
    pdf.add_font("Noto", "B", bold_font)

    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Title
    report_type = data.get("report_type", "تقرير")
    title_map = {
        "monthly": "شهري",
        "quarterly": "ربع سنوي",
        "annual": "سنوي",
        "ncnp": "لـ NCNP",
    }
    title = title_map.get(report_type, report_type)
    period = data.get("period", "")
    pdf.set_font("Noto", "B", 16)
    pdf.set_text_color(0, 51, 102)
    pdf.cell(0, 10, _to_display_text(f"تقرير {title}"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Noto", "", 11)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(0, 8, _to_display_text(f"الفترة: {period}"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 8, _to_display_text(f"تاريخ التوليد: {data.get('generated_at', '')}"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)

    # Finance summary
    finance = data.get("finance", {})
    pdf.set_font("Noto", "B", 13)
    pdf.cell(0, 10, _to_display_text("ملخص الحوكمة والمالية"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Noto", "", 11)
    pdf.cell(0, 8, _to_display_text(f"إجمالي الإيرادات: {finance.get('total_income', 0):,.2f} ريال"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 8, _to_display_text(f"إجمالي المصروفات: {finance.get('total_expenses', 0):,.2f} ريال"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 8, _to_display_text(f"الرصيد: {finance.get('balance', 0):,.2f} ريال"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)

    # Programs
    programs = data.get("programs", [])
    pdf.set_font("Noto", "B", 13)
    pdf.cell(0, 10, _to_display_text("حالة البرامج"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Noto", "", 11)
    for p in programs:
        line = (
            f"{p.get('name', '')} - الحالة: {p.get('status', '')} - "
            f"التقدم الكلي: {p.get('overall_progress', 0)}٪"
        )
        pdf.cell(0, 8, _to_display_text(line), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)

    # Volunteers
    volunteer = data.get("volunteers", {})
    pdf.set_font("Noto", "B", 13)
    pdf.cell(0, 10, _to_display_text("ملخص التطوع"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Noto", "", 11)
    pdf.cell(0, 8, _to_display_text(f"إجمالي المتطوعين: {volunteer.get('total_volunteers', 0)}"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 8, _to_display_text(f"الفرص المفتوحة: {volunteer.get('open_opportunities', 0)}"), align="R", new_x="LMARGIN", new_y="NEXT")

    pdf.output(str(output_path))
    return str(output_path)


def export_word(data: dict[str, Any], base_name: str) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    report_dir = settings.reports_dir
    report_dir.mkdir(parents=True, exist_ok=True)
    output_path = report_dir / f"{base_name}.docx"

    doc = Document()
    doc.styles["Normal"].font.name = "Noto Sans Arabic"
    doc.styles["Normal"].font.size = None

    def add_rtl_paragraph(text: str, bold: bool = False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.rtl = True
        run = p.add_run(text)
        run.font.rtl = True
        run.font.bold = bold
        return p

    add_rtl_paragraph("منصة وكلاء الذكاء الاصطناعي للقطاع غير الربحي", bold=True)
    add_rtl_paragraph(f"تقرير {data.get('report_type', '')}", bold=True)
    add_rtl_paragraph(f"الفترة: {data.get('period', '')}")
    add_rtl_paragraph(f"تاريخ التوليد: {data.get('generated_at', '')}")

    add_rtl_paragraph("ملخص الحوكمة والمالية", bold=True)
    finance = data.get("finance", {})
    add_rtl_paragraph(f"إجمالي الإيرادات: {finance.get('total_income', 0):,.2f} ريال")
    add_rtl_paragraph(f"إجمالي المصروفات: {finance.get('total_expenses', 0):,.2f} ريال")
    add_rtl_paragraph(f"الرصيد: {finance.get('balance', 0):,.2f} ريال")

    add_rtl_paragraph("حالة البرامج", bold=True)
    for p in data.get("programs", []):
        line = (
            f"{p.get('name', '')} - الحالة: {p.get('status', '')} - "
            f"التقدم الكلي: {p.get('overall_progress', 0)}٪"
        )
        add_rtl_paragraph(line)

    add_rtl_paragraph("ملخص التطوع", bold=True)
    volunteer = data.get("volunteers", {})
    add_rtl_paragraph(f"إجمالي المتطوعين: {volunteer.get('total_volunteers', 0)}")
    add_rtl_paragraph(f"الفرص المفتوحة: {volunteer.get('open_opportunities', 0)}")

    doc.save(str(output_path))
    return str(output_path)


def export_excel(data: dict[str, Any], base_name: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    report_dir = settings.reports_dir
    report_dir.mkdir(parents=True, exist_ok=True)
    output_path = report_dir / f"{base_name}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "تقرير"
    ws.sheet_view.rightToLeft = True

    def write_row(row: int, values: list[Any], bold: bool = False):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(horizontal="right")
            cell.font = Font(name="Noto Sans Arabic", bold=bold)

    write_row(1, ["تقرير", data.get("report_type", "")], bold=True)
    write_row(2, ["الفترة", data.get("period", "")])
    write_row(3, ["تاريخ التوليد", data.get("generated_at", "")])

    finance = data.get("finance", {})
    write_row(5, ["ملخص الحوكمة والمالية"], bold=True)
    write_row(6, ["إجمالي الإيرادات", finance.get("total_income", 0)])
    write_row(7, ["إجمالي المصروفات", finance.get("total_expenses", 0)])
    write_row(8, ["الرصيد", finance.get("balance", 0)])

    write_row(10, ["البرامج", "الحالة", "التقدم الكلي"], bold=True)
    for i, p in enumerate(data.get("programs", []), start=11):
        write_row(i, [p.get("name", ""), p.get("status", ""), p.get("overall_progress", 0)])

    volunteer = data.get("volunteers", {})
    start = 11 + len(data.get("programs", [])) + 1
    write_row(start, ["ملخص التطوع"], bold=True)
    write_row(start + 1, ["إجمالي المتطوعين", volunteer.get("total_volunteers", 0)])
    write_row(start + 2, ["الفرص المفتوحة", volunteer.get("open_opportunities", 0)])

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 20

    wb.save(str(output_path))
    return str(output_path)
